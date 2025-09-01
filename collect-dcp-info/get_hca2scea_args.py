import os
import signal
import argparse
import requests
import datetime
import pandas as pd
from tqdm import tqdm
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from hca_ingest.api.ingestapi import IngestApi
from hca_ingest.downloader.workbook import WorkbookDownloader

# Constants
INGEST_API_URL = "https://api.ingest.archive.data.humancellatlas.org"
SCEA_ARGS_TEMPLATE = {
    "-s": None, "-id": None, "-c": None, "-ac": None, "-et": None,
    "-f": None, "-pd": None, "-hd": None, "-study": None,
    "--facs": False, "-name": None, "-o": None
}

pd.set_option('display.max_colwidth', None)  

# Timeout handler
class InputTimedOut(Exception):
    pass

def inputTimeOutHandler(signum, frame):
    # called when read times out
    print('time\'s up!')
    raise InputTimedOut

signal.signal(signal.SIGALRM, inputTimeOutHandler)

def input_with_timeout(prompt_string, timeout=0, default_value=None):
    foo = default_value
    try:
        print(prompt_string, end=' ')
        if timeout > 0:
            signal.alarm(timeout)
        foo = input()
        signal.alarm(0)    #disable alarm
    except InputTimedOut:
            pass
    return foo

def get_valid_api(token=None):
    api = IngestApi(INGEST_API_URL)
    api.set_token(f"Bearer {token}")
    response = requests.get(f"{INGEST_API_URL}/submissionEnvelopes/", headers=api.get_headers())
    while response.status_code != 200:
        token = input_with_timeout("Please provide a valid token:", timeout=10, default_value=token).strip()
        api.set_token(f"Bearer {token}")
        response = requests.get(f"{INGEST_API_URL}/submissionEnvelopes/", headers=api.get_headers())
    return api

def get_valid_submission_uuid(api: IngestApi):
    resp = requests.get(f"{INGEST_API_URL}/submissionEnvelopes/search/findByUuidUuid?uuid={uuid}", headers=api.get_headers())
    while True:
        uuid = input_with_timeout("Please provide a valid submission uuid:", timeout=10, default_value=uuid).strip()
        resp = requests.get(f"{INGEST_API_URL}/submissionEnvelopes/search/findByUuidUuid?uuid={uuid}", headers=api.get_headers())
        if resp.status_code == 200:
            return uuid
        print("Invalid submission UUID.")

def download_workbook(api: IngestApi, sub_uuid: str):
    submission = api.get_submission_by_uuid(sub_uuid)
    project = api.get_related_project(submission['_links']['self']['href'].split('/')[-1])
    if project:
        project_title = project['content']['project_core']['project_short_name']
        if os.path.exists(f"hca_spreadsheets/{project_title}.xlsx"):
            print(f"Workbook for {project_title} already exists. Skipping download.", flush=True)
            return load_workbook(f"hca_spreadsheets/{project_title}.xlsx")
    print("Downloading workbook...", flush=True, end=' ')
    try:
        wd = WorkbookDownloader(api)
        return wd.get_workbook_from_submission(sub_uuid)
    except Exception as e:
        print(f"Error downloading workbook: {e}", flush=True)
        return None

def save_workbook(wb, proj_sheet='Project'):
    proj_name_cell = [col[3].column_letter + "6" for col in wb[proj_sheet].iter_cols()
                      if col[3].value == 'project.project_core.project_short_name']
    proj_name = wb[proj_sheet][proj_name_cell[0]].value
    file_name = f"{proj_name}.xlsx"
    wb.save(f"hca_spreadsheets/{file_name}")
    print(f"Saving workbook as {file_name}.", flush=True)
    return file_name, proj_name

def extract_project_metadata(xl: pd.ExcelFile, scea_args: dict):
    print("Extracting: uuid", flush=True, end=' ')
    project = xl.parse('Project', header=None)
    scea_args['-id'] = project.loc[5, project.iloc[3] == 'project.uuid'].values[0]
    
    # Study accession
    print(", insdc_project_accessions", flush=True, end=' ')
    insdc_vals = project.loc[5, project.iloc[3] == 'project.insdc_project_accessions']
    if insdc_vals.notna().all() and not insdc_vals.empty:
        acc_str = insdc_vals.values[0]
        if '||' in acc_str:
            acc_list = acc_str.split('||')
            print("\nMultiple INSDC accessions found:", flush=True)
            acc_list_df = pd.DataFrame({'accession': acc_list})
            acc_list_df['study_name'] = acc_list_df['accession'].apply(fetch_ena_study_name)
            print(acc_list_df, flush=True)
            index = int(input_with_timeout("Select an accession: ", timeout=10, default_value=0))
            scea_args['-study'] = acc_list[index]
        elif acc_str.startswith(('ERP', 'DRP', 'SRP')):
            scea_args['-study'] = acc_str

def fetch_hca_update_date(proj_uuid: str):
    print("Getting hca update date...", flush=True)
    url = f"{INGEST_API_URL}/projects/search/findByUuid?uuid={proj_uuid}"
    hd = requests.get(url).json()['updateDate']
    return datetime.datetime.fromisoformat(hd).strftime("%Y-%m-%d")

def fetch_ena_publication_date(study_accession: str, ena_value="ENA-FIRST-PUBLIC"):
    print(f"Getting {ena_value}...", flush=True)
    url = f"https://www.ebi.ac.uk/ena/browser/api/xml/{study_accession}"
    root = ET.fromstring(requests.get(url).text)
    for attr in root.findall('.//STUDY_ATTRIBUTE'):
        if attr.find('TAG').text == ena_value:
            return attr.find('VALUE').text
    return None

def fetch_ena_study_name(study_accession: str):
    url = f"https://www.ebi.ac.uk/ena/browser/api/xml/{study_accession}"
    root = ET.fromstring(requests.get(url).text)
    for attr in root.findall('.//DESCRIPTOR'):
        if attr.find('STUDY_TITLE') is not None:
            return attr.find('STUDY_TITLE').text
    return None

def detect_experiment_type_and_factor(xl, factors_map):
    print("Extracting experiment type and factor...", flush=True, end=' ')
    for factor, col_name in factors_map.items():
        sheet_name = col_name.split('.')[0].replace('_', ' ').capitalize()
        try:
            df = xl.parse(sheet_name, header=None)
        except ValueError:
            continue

        col_idx = next((i for i, val in enumerate(df.iloc[3]) if val == col_name), None)
        if col_idx is None:
            continue

        values = df.iloc[5:, col_idx].dropna().unique()
        if len(values) > 1 and factor != 'individual':
            print(f"Selecting factor [{factor}] with values: {values}", flush=True)
            return 'differential', factor
    return 'baseline', 'individual'

def choose_name_field(xl):
    print(f"Selecting sample name.", flush=True)
    if 'Cell suspension' not in xl.sheet_names:
        return 'sp_id'
    cs_sheet = xl.parse("Cell suspension", header=None)
    sp_sheet = xl.parse("Specimen from organism", header=None)

    cs_id_count = cs_sheet.loc[5:, cs_sheet.iloc[3] == 'cell_suspension.biomaterial_core.biomaterial_id'].shape[0]
    sp_id_count = sp_sheet.loc[5:, sp_sheet.iloc[3] == 'specimen_from_organism.biomaterial_core.biomaterial_id'].shape[0]

    return 'cs_id' if cs_id_count >= sp_id_count else 'sp_id'

def check_facs_used(xl):
    print("Checking if FACS was used... ", flush=True, end=' ')
    if 'Enrichment protocol' in xl.sheet_names:
        enrich = xl.parse('Enrichment protocol', header=None)
        ontology_col = enrich.iloc[2] == 'enrichment_protocol.method.ontology'
        if ontology_col.any():
            used = enrich.loc[4:, ontology_col].isin(['EFO:0009108']).any().any()
            print("yes" if used else "no")
            return used
        print("no")
    return False

def append_args_to_csv(scea_args: dict, submission_uuid: str, csv_path: str):
    row_data = {'submission_uuid': submission_uuid}
    row_data.update(scea_args)

    df_row = pd.DataFrame([row_data])

    if os.path.exists(csv_path):
        df_csv = pd.read_csv(csv_path)
        df_csv = pd.concat([df_csv, df_row], ignore_index=True)
    else:
        df_csv = df_row

    df_csv.to_csv(csv_path, index=False)
    print(f"✅ Appended submission {submission_uuid} to {csv_path}")

def filter_scea_submission(df):
    print("Filtering submissions for SCEA...")
    df = df[df['azul_valid']]
    df = df[~df['cxg_geo'] & ~df['cxg_link'] & (df['cxg_doi'] == "False")]
    df = df[~df['lib_prots'].str.contains('\\|\\|', na=True)]
    df = df[~df['organisms'].str.contains('\\|\\|', na=True)]
    df = df[df['analysis_types'].str.contains('type|annotation|sample', na=False)]
    return df

def main():
    parser = argparse.ArgumentParser(description="Process HCA submissions for SCEA.")
    parser.add_argument("-u", "--uuid", type=str, help="Single submission UUID to process.")
    parser.add_argument("-f", "--csv", type=str, default="hca_submissions_summary.csv", help="CSV file with submission UUIDs.")
    parser.add_argument("-t", "--token", type=str, help="Bearer token for HCA ingest API.")
    parser.add_argument("-c", "--curator", type=str, required=True, help="Curator initials (capitalized).")
    args = parser.parse_args()

    curator = args.curator.strip().upper()
    api = get_valid_api(args.token)

    if args.uuid:
        print(f"Processing single submission UUID: {args.uuid.strip()}")
        submission_uuids = [args.uuid.strip()]
    else:
        if not os.path.exists(args.csv):
            if not os.path.exists(os.path.join("collect-dcp-info", args.csv)):
                raise FileNotFoundError(f"CSV file {args.csv} not found.")
            else:
                args.csv = os.path.join("collect-dcp-info", args.csv)
        print(f"Processing multiple submissions from CSV file: {args.csv}")
        submission_df = pd.read_csv(args.csv)
        submission_df['organisms'] = submission_df['organisms'].astype(str)
        submission_df_filtered = filter_scea_submission(submission_df)
        submission_uuids = submission_df_filtered['sub_uuid'].dropna().unique()
        print(f"Found {submission_df_filtered.shape[0]}/{submission_df.shape[0]} submissions eligible for scea in the {args.csv}.")
        if submission_uuids.size == 0:
            return

    for sub_uuid in tqdm(submission_uuids, desc="Processing submissions", unit="submission"):
        print(f"Getting submission {sub_uuid}")
        wb = download_workbook(api, sub_uuid)
        if wb is None:
            print(f"Skipping submission {sub_uuid} due to download error.", flush=True)
            continue
        file_name, proj_name = save_workbook(wb)
        xl = pd.ExcelFile(f"hca_spreadsheets/{file_name}")

        scea_args = SCEA_ARGS_TEMPLATE.copy()
        scea_args['-s'] = os.path.join("hca_spreadsheets", file_name)
        scea_args['-o'] = os.path.join("hca_spreadsheets", proj_name)
        os.makedirs(scea_args['-o'], exist_ok=True)

        extract_project_metadata(xl, scea_args)
        scea_args['-c'] = curator
        scea_args['-hd'] = fetch_hca_update_date(scea_args['-id'])

        if scea_args.get('-study'):
            scea_args['-pd'] = fetch_ena_publication_date(scea_args['-study'])

        factors = {
            'disease': 'donor_organism.diseases.text',
            'development_stage': 'donor_organism.development_stage.text',
            'organ': 'specimen_from_organism.organ.text',
            'individual': 'donor_organism.biomaterial_core.biomaterial_id'
        }
        scea_args['-et'], scea_args['-f'] = detect_experiment_type_and_factor(xl, factors)
        scea_args['-name'] = choose_name_field(xl)
        scea_args['--facs'] = check_facs_used(xl)

        append_args_to_csv(scea_args, sub_uuid, 'scea_arguments.csv')

        print("\n🎯Submission processed.")
        print("📄 Output written to: scea_arguments.csv")

if __name__ == "__main__":
    main()
